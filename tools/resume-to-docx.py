"""Build an editable .docx from resume-src.html so the Google Docs copy can be
re-imported without losing the corrections. Parses the same HTML the PDF is
rendered from, through the same loader, so the two cannot drift apart.

    python tools/resume-to-docx.py ~/Desktop/Nicolas_Ruth_Resume.docx

This produces the FULL variant -- it carries the phone number and is meant for
direct applications, not for publishing. It therefore needs the gitignored
tools/resume-private.json; see resume_source.py.
"""
import re, html, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import resume_source

if len(sys.argv) < 2:
    sys.exit(__doc__.strip())
OUT = sys.argv[-1]
if not OUT.lower().endswith('.docx'):
    sys.exit(f"expected a .docx output path, got {OUT!r}")
try:
    raw = resume_source.load("full")
except resume_source.MissingPrivateFile as exc:
    sys.exit(str(exc))
body = re.search(r'<body>(.*)</body>', raw, re.S).group(1)

# ---------- inline runs: (text, bold, italic) ----------
def runs(frag):
    """Split an HTML fragment into styled runs. <b> -> bold, .yr/.org spans -> italic/plain."""
    frag = re.sub(r'<span class="org">(.*?)</span>', r'\1', frag, flags=re.S)
    out, pos = [], 0
    pat = re.compile(r'<b>(.*?)</b>|<span class="yr">(.*?)</span>', re.S)
    for m in pat.finditer(frag):
        if m.start() > pos:
            out.append((frag[pos:m.start()], False, False))
        if m.group(1) is not None:
            out.append((m.group(1), True, False))
        else:
            out.append((m.group(2), False, True))
        pos = m.end()
    out.append((frag[pos:], False, False))
    clean = []
    for t, b, i in out:
        t = re.sub(r'<[^>]+>', '', t)
        t = html.unescape(t)
        t = re.sub(r'\s+', ' ', t)
        if t:
            clean.append((t, b, i))
    return clean

def plain(frag):
    return ''.join(t for t, _, _ in runs(frag))

# ---------- document setup ----------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.40)
sec.left_margin = sec.right_margin = Inches(0.50)
RIGHT_TAB = Inches(8.5 - 0.50 - 0.50)

st = doc.styles['Normal']
st.font.name = 'Arial'
st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
pf = st.paragraph_format
pf.space_before = pf.space_after = Pt(0)
pf.line_spacing = 1.0

def para(space_before=0, space_after=0, left=0, hang=0, align=None):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(space_before), Pt(space_after)
    if left:
        f.left_indent = Pt(left)
    if hang:
        f.first_line_indent = Pt(-hang)
    if align is not None:
        p.alignment = align
    return p

def emit(p, frag, bold_all=False):
    for t, b, i in runs(frag):
        r = p.add_run(t)
        r.bold = b or bold_all
        r.italic = i
    return p

def rule(p):
    """Bottom border on a paragraph -- the section underline."""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    bdr.append(bottom)
    pPr.append(bdr)

# ---------- walk the blocks in document order ----------
# NOTE: .entry contains nested <div>s (the entry-head), so a non-greedy regex
# closes on the wrong tag. Scan and count div depth instead.
def div_body(s, open_at):
    """Given the index of a '<div', return (inner_html, index_after_closing_div)."""
    depth, i = 0, open_at
    start = None
    while i < len(s):
        if s.startswith('<div', i):
            depth += 1
            if start is None:
                start = s.index('>', i) + 1
            i = s.index('>', i) + 1
        elif s.startswith('</div>', i):
            depth -= 1
            if depth == 0:
                return s[start:i], i + 6
            i += 6
        else:
            i += 1
    raise ValueError('unbalanced div')

def blocks(s):
    """Yield (kind, payload) in document order."""
    starts = [
        ('name',    re.compile(r'<p class="name">')),
        ('contact', re.compile(r'<p class="contact">')),
        ('h2',      re.compile(r'<h2>')),
        ('oneline', re.compile(r'<p class="oneline">')),
        ('entry',   re.compile(r'<div class="entry">')),
        ('skills',  re.compile(r'<div class="skills">')),
    ]
    i = 0
    while i < len(s):
        best = None
        for kind, pat in starts:
            mm = pat.search(s, i)
            if mm and (best is None or mm.start() < best[1].start()):
                best = (kind, mm)
        if best is None:
            return
        kind, mm = best
        if kind in ('entry', 'skills'):
            inner, nxt = div_body(s, mm.start())
            yield kind, inner
            i = nxt
        else:
            close = '</h2>' if kind == 'h2' else '</p>'
            end = s.index(close, mm.end())
            yield kind, s[mm.end():end]
            i = end + len(close)

class M:
    def __init__(self, kind, payload):
        self.kind, self.payload = kind, payload
    def group(self, k):
        return self.payload if k == self.kind else None

for m in (M(k, v) for k, v in blocks(body)):
    if m.group('name'):
        p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(plain(m.group('name')))
        r.bold = True
        r.font.size = Pt(18)

    elif m.group('contact'):
        p = para(space_before=1, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(plain(m.group('contact')))
        r.font.size = Pt(9.3)

    elif m.group('h2'):
        p = para(space_before=5, space_after=1.5)
        r = p.add_run(plain(m.group('h2')).upper())
        r.bold = True
        r.font.size = Pt(9.8)
        rule(p)

    elif m.group('entry'):
        e = m.group('entry')
        head = re.search(r'<div class="entry-head">(.*?)</div>', e, re.S)
        if head:
            h = head.group(1)
            title = re.search(r'<span class="entry-title">(.*?)</span>\s*</span>|'
                              r'<span class="entry-title">(.*?)</span>', h, re.S)
            tfrag = re.search(r'<span class="entry-title">(.*)</span>\s*<span class="entry-meta',
                              h, re.S)
            meta = re.search(r'<span class="entry-meta[^"]*">(.*?)</span>', h, re.S)
            p = para(space_before=3.5)
            p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
            emit(p, tfrag.group(1) if tfrag else h)
            if meta:
                p.add_run('\t')
                r = p.add_run(plain(meta.group(1)))
                r.italic = 'tools' not in head.group(1)[:200] and 'entry-meta tools' not in h
                r.font.size = Pt(9.1)
        # the degree line sits between the head and the bullets
        deg = re.search(r'<div class="degree">(.*?)</div>', e, re.S)
        if deg:
            emit(para(space_before=0.6), deg.group(1))
        for li in re.findall(r'<li>(.*?)</li>', e, re.S):
            p = para(space_before=0.8, left=10, hang=10)
            p.add_run('•  ')
            emit(p, li)

    elif m.group('oneline'):
        p = para(space_before=3)
        emit(p, m.group('oneline'))

    elif m.group('skills'):
        for sp in re.findall(r'<p>(.*?)</p>', m.group('skills'), re.S):
            p = para(space_before=0.8, left=44, hang=44)
            emit(p, sp)

doc.save(OUT)
print('wrote', OUT)
